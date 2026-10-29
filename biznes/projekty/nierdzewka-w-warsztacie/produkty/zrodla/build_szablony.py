#!/usr/bin/env python3
"""Generuje 8 szablonów dokumentacji spawalniczej (puste + przykłady wypełnione).

Wymaga: python-docx, openpyxl. Uruchom: python3 build_szablony.py
Wynik: ../gotowe/szablony/puste/*.docx|xlsx i ../gotowe/szablony/przyklady/*.docx|xlsx
Treść: własny układ pól; bez kopiowania treści norm (tylko numery norm jako odwołania).
Dane w przykładach są FIKCYJNE.
"""
import os
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BAZA = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BAZA, "..", "gotowe", "szablony")
AKCENT = RGBColor(0x0B, 0x6F, 0xB8)
SZARY = "EEF3F8"
WERSJA = "Szablony warsztatu v1.0 · Instal-tig / Nierdzewka w warsztacie"
FIRMA_P = "Warsztat Przykładowy (DANE FIKCYJNE), ul. Przykładowa 1, 00-000 Miasto"


# ------------------------------------------------------------------ DOCX - pomocnicze
def cieniuj(cell, kolor=SZARY):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), kolor)
    tcPr.append(shd)


def szerokosci(t, mm):
    t.autofit = False
    for row in t.rows:
        for c, w in zip(row.cells, mm):
            c.width = Mm(w)


def nowy_doc(tytul, nr_pola, przyklad, poziomo=False):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    if poziomo:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Mm(297), Mm(210)
    for m in ("left_margin", "right_margin"):
        setattr(s, m, Mm(15))
    s.top_margin, s.bottom_margin = Mm(14), Mm(14)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9.5)
    # nagłówek: logo/firma + tytuł
    h = s.header.paragraphs[0]
    h.text = (FIRMA_P if przyklad else "[LOGO I DANE FIRMY - wstaw tutaj]") + "\t\t" + (
        "PRZYKŁAD WYPEŁNIENIA - dane fikcyjne" if przyklad else "")
    h.runs[0].font.size = Pt(8)
    f = s.footer.paragraphs[0]
    f.text = WERSJA + " · Szablon nie zastępuje wymagań zamawiającego, WPQR ani uprawnień."
    f.runs[0].font.size = Pt(7)
    f.runs[0].font.color.rgb = RGBColor(0x5B, 0x67, 0x76)
    p = d.add_paragraph()
    r = p.add_run(tytul)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = AKCENT
    if nr_pola:
        tabela_kv(d, nr_pola)
    return d


def tabela_kv(d, pary, szer=(60, 120)):
    """Tabela etykieta | wartość (2 kolumny) lub 4 kolumny gdy pary w parach."""
    t = d.add_table(rows=0, cols=4 if isinstance(pary[0][0], tuple) else 2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for para in pary:
        cells = t.add_row().cells
        if isinstance(para[0], tuple):
            for i, (k, v) in enumerate(para):
                cells[2 * i].text, cells[2 * i + 1].text = k, v
                cieniuj(cells[2 * i])
                cells[2 * i].paragraphs[0].runs[0].bold = True
        else:
            k, v = para
            cells[0].text, cells[1].text = k, v
            cieniuj(cells[0])
            cells[0].paragraphs[0].runs[0].bold = True
    szerokosci(t, [40, 50, 40, 50] if len(t.columns) == 4 else list(szer))
    d.add_paragraph()
    return t


def tabela(d, naglowki, wiersze, puste=0, szer=None):
    t = d.add_table(rows=1, cols=len(naglowki))
    t.style = "Table Grid"
    for i, n in enumerate(naglowki):
        c = t.rows[0].cells[i]
        c.text = n
        cieniuj(c)
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].runs[0].font.size = Pt(8.5)
    for w in wiersze:
        cells = t.add_row().cells
        for i, v in enumerate(w):
            cells[i].text = str(v)
    for _ in range(puste):
        t.add_row()
    szerokosci(t, szer or [180 / len(naglowki)] * len(naglowki))
    d.add_paragraph()
    return t


def sekcja(d, tekst):
    p = d.add_paragraph()
    r = p.add_run(tekst)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = AKCENT


def uwaga(d, tekst):
    p = d.add_paragraph()
    r = p.add_run(tekst)
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x5B, 0x67, 0x76)


def podpisy(d, role):
    t = d.add_table(rows=2, cols=len(role))
    t.style = "Table Grid"
    for i, rr in enumerate(role):
        t.rows[0].cells[i].text = rr
        cieniuj(t.rows[0].cells[i])
        t.rows[1].cells[i].text = "\n\nimię, nazwisko, data, podpis"
    szerokosci(t, [180 / len(role)] * len(role))
    d.add_paragraph()


def v(przyklad, wartosc):
    return wartosc if przyklad else ""


# ------------------------------------------------------------------ 1. WPS
def wps(przyklad):
    P = przyklad
    d = nowy_doc("Instrukcja technologiczna spawania (WPS) - TIG 141", [
        (("Nr WPS", v(P, "WPS-141-01")), ("Rewizja / data", v(P, "0 / 2026-11-05"))),
        (("Wytwórca", v(P, "Warsztat Przykładowy (fikcyjny)")), ("Miejsce spawania", v(P, "warsztat i montaż"))),
        (("Podstawa kwalifikacji (nr WPQR / inne)", v(P, "WPQR nr [[nr z jednostki]]")), ("Norma odniesienia układu", "PN-EN ISO 15609-1")),
    ], P)
    uwaga(d, "WPS bez podstawy kwalifikacji (np. WPQR) nie jest kwalifikowaną technologią. Zakres stosowania wynika z WPQR i wymagań zamawiającego.")
    sekcja(d, "1. Złącze")
    tabela_kv(d, [
        (("Rodzaj złącza / spoiny", v(P, "doczołowe, rura-rura, BW")), ("Przygotowanie", v(P, "na styk, bez ukosowania"))),
        (("Szczelina", v(P, "0-0,5 mm")), ("Niewspółosiowość maks.", v(P, "wg wymagań klienta"))),
    ])
    p = d.add_paragraph("Szkic złącza i kolejność ściegów:")
    t = d.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = v(P, "rura Ø 51 × 2,0 mm, styk, 1 ścieg (przetop + lico) [[wstaw szkic]]") + "\n\n\n\n\n"
    szerokosci(t, [180])
    d.add_paragraph()
    sekcja(d, "2. Materiał podstawowy i zakresy")
    tabela_kv(d, [
        (("Gatunek (nr / nazwa)", v(P, "1.4404 / 316L")), ("Grupa materiałowa", v(P, "[[wg WPQR]]"))),
        (("Zakres grubości ścianki", v(P, "[[wg WPQR]]")), ("Zakres średnic zewn.", v(P, "[[wg WPQR]]"))),
    ])
    sekcja(d, "3. Proces, pozycje, materiały dodatkowe")
    tabela_kv(d, [
        (("Proces", "141 (TIG)"), ("Pozycje spawania", v(P, "PA, PC, PH"))),
        (("Technika", v(P, "prosta")), ("Spoiwo (oznaczenie, Ø)", v(P, "pręt 316L, Ø 1,6 mm"))),
        (("Gaz osłonowy / przepływ", v(P, "Ar 4.8 / 7-9 l/min")), ("Gaz formujący / przepływ / przedmuch", v(P, "Ar 4.8 / 4-6 l/min / 2 min"))),
        (("Tlen resztkowy dopuszczalny", v(P, "≤ 50 ppm (wym. klienta) [[DO SPRAWDZENIA]]")), ("Wolfram (rodzaj, Ø, kąt)", v(P, "WL20, 1,6 mm, ok. 30°"))),
    ])
    sekcja(d, "4. Parametry spawania")
    tabela(d, ["Ścieg", "Ø spoiwa [mm]", "Prąd [A]", "Napięcie [V]", "Rodzaj prądu / biegun.", "Prędkość [cm/min]", "Energia liniowa [kJ/mm]"],
           [["1", "1,6", "65-80", "9-11", "DC (-)", "6-9", "[[wyliczyć]]"]] if P else [], puste=0 if P else 4)
    sekcja(d, "5. Ciepło i obróbka po spawaniu")
    tabela_kv(d, [
        (("Podgrzewanie wstępne", v(P, "nie dotyczy")), ("Maks. temp. międzyściegowa", v(P, "150°C [[wg WPQR]]"))),
        (("Czyszczenie / trawienie / pasywacja", v(P, "trawienie żelem, spłukanie, pasywacja")), ("Inne", v(P, "-"))),
    ])
    podpisy(d, ["Sporządził", "Zatwierdził (koordynator spawania)"])
    return d


# ------------------------------------------------------------------ 3. VT
def vt(przyklad):
    P = przyklad
    d = nowy_doc("Protokół badań wizualnych złączy spawanych (VT)", [
        (("Nr protokołu", v(P, "VT/2026/014")), ("Data", v(P, "2026-11-12"))),
        (("Zlecenie / obiekt", v(P, "Linia CIP - rozbudowa (fikcyjne)")), ("Zamawiający", v(P, "Mleczarnia Przykładowa (fikcyjna)"))),
        (("Nr rysunku / izometrii", v(P, "ISO-CIP-03 rew. B")), ("Etap", v(P, "po spawaniu, przed trawieniem"))),
    ], P)
    sekcja(d, "Warunki i kryteria")
    tabela_kv(d, [
        (("Badanie wg", "PN-EN ISO 17637"), ("Kryteria akceptacji wg", v(P, "PN-EN ISO 5817"))),
        (("Poziom jakości", v(P, "B")), ("Wymagania dodatkowe klienta", v(P, "grań: kolor maks. poz. 2 (AWS D18.2)"))),
        (("Oświetlenie [lx]", v(P, "≥ 500 [[DO SPRAWDZENIA - wymaganie]]")), ("Dostęp", v(P, "lico bezpośredni, grań endoskop"))),
        (("Narzędzia", v(P, "lupa, spoinomierz, endoskop")), ("Stan powierzchni", v(P, "czysta, bez odprysków"))),
    ])
    sekcja(d, "Wyniki")
    tabela(d, ["Nr złącza", "Wymiar", "WPS", "Spawacz", "Lico", "Grań", "Niezgodności (opis)", "Ocena"],
           [["S-01", "Ø51×2", "WPS-141-01", "AB", "OK", "OK, poz. 1", "-", "OK"],
            ["S-02", "Ø51×2", "WPS-141-01", "AB", "OK", "poz. 3", "nalot w grani ponad wymaganie", "nOK - do wycięcia"],
            ["S-03", "Ø51×2", "WPS-141-01", "CD", "podtopienie 0,2 mm", "OK, poz. 1", "podtopienie lokalne", "OK (poziom B) [[DO SPRAWDZENIA]]"]] if P else [],
           puste=0 if P else 12, szer=[16, 18, 22, 16, 22, 22, 40, 24])
    tabela_kv(d, [("Wynik ogólny", v(P, "2 złącza OK, 1 do naprawy - patrz S-02")), ("Uwagi", v(P, "S-02 wycięte i spawane ponownie jako S-02A - osobny protokół"))])
    podpisy(d, ["Badający (kwalifikacje)", "Przedstawiciel zamawiającego"])
    return d


# ------------------------------------------------------------------ 4. Próba szczelności
def proba(przyklad):
    P = przyklad
    d = nowy_doc("Protokół próby szczelności / ciśnieniowej", [
        (("Nr protokołu", v(P, "PC/2026/006")), ("Data", v(P, "2026-11-14"))),
        (("Obiekt / odcinek / nr linii", v(P, "Linia CIP, odcinek L3")), ("Zamawiający", v(P, "Mleczarnia Przykładowa (fikcyjna)"))),
    ], P)
    uwaga(d, "Próby medium ściśliwym (powietrze, azot) gromadzą energię - stosuj procedury BHP, wymagania zamawiającego i przepisy (np. UDT), jeśli dotyczą.")
    sekcja(d, "Parametry próby")
    tabela_kv(d, [
        (("Medium próby", v(P, "woda")), ("Temperatura medium", v(P, "12°C"))),
        (("Ciśnienie robocze", v(P, "6 bar")), ("Ciśnienie próby", v(P, "9 bar [[wg wymagań klienta]]"))),
        (("Czas utrzymania", v(P, "30 min")), ("Dopuszczalny spadek", v(P, "0 bar"))),
        (("Manometr (nr, zakres, klasa)", v(P, "M-12, 0-16 bar, kl. 1,0")), ("Wzorcowanie / legalizacja do", v(P, "2027-03-31"))),
    ])
    sekcja(d, "Przebieg")
    tabela(d, ["Godzina", "Ciśnienie [bar]", "Uwagi"],
           [["10:00", "9,0", "osiągnięcie ciśnienia próby"], ["10:15", "9,0", "oględziny złączy - brak wycieków"], ["10:30", "9,0", "koniec próby"]] if P else [],
           puste=0 if P else 8, szer=[30, 40, 110])
    tabela_kv(d, [("Nieszczelności (miejsce)", v(P, "brak")), ("Wynik próby", v(P, "POZYTYWNY"))])
    podpisy(d, ["Wykonawca", "Przedstawiciel zamawiającego"])
    return d


# ------------------------------------------------------------------ 5. Odbiór
def odbior(przyklad):
    P = przyklad
    d = nowy_doc("Protokół odbioru i przekazania instalacji", [
        ("Nr protokołu / data", v(P, "OD/2026/003 / 2026-11-20")),
        ("Wykonawca", v(P, FIRMA_P)),
        ("Zamawiający", v(P, "Mleczarnia Przykładowa (fikcyjna), ul. Mleczna 2, 00-000 Miasto")),
        ("Umowa / zamówienie nr", v(P, "ZAM/2026/118")),
        ("Przedmiot i zakres odbioru", v(P, "Rurociągi CIP L1-L3, 316L, Ø51-Ø76, ok. 64 m")),
    ], P)
    sekcja(d, "Przekazana dokumentacja")
    docs = ["Instrukcje technologiczne spawania (WPS) + podstawa kwalifikacji", "Dziennik spawania", "Protokoły badań wizualnych (VT)",
            "Protokoły innych badań (PT / RT)", "Protokoły prób szczelności / ciśnieniowych", "Rejestr atestów materiałowych + kopie świadectw",
            "Uprawnienia spawaczy (kopie)", "Rysunki / izometrie powykonawcze", "Instrukcje / karty urządzeń", "Inne:"]
    zazn = {0, 1, 2, 4, 5, 6, 7}
    tabela(d, ["Dokument", "Przekazano (TAK / NIE / ND)", "Liczba stron / uwagi"],
           [[x, ("TAK" if i in zazn else "ND") if P else "", ""] for i, x in enumerate(docs)], szer=[95, 40, 45])
    sekcja(d, "Stwierdzone usterki")
    tabela(d, ["Lp.", "Opis usterki", "Termin usunięcia", "Odpowiedzialny"],
           [["1", "brak oznakowania kierunku przepływu na L2", "2026-11-27", "wykonawca"]] if P else [], puste=0 if P else 4, szer=[12, 98, 35, 35])
    tabela_kv(d, [("Decyzja", v(P, "odbiór z uwagami (usterka nr 1)") or "☐ odbiór bez uwag   ☐ odbiór z uwagami   ☐ odmowa odbioru"),
                  ("Data rozpoczęcia gwarancji", v(P, "2026-11-20"))])
    podpisy(d, ["Za wykonawcę", "Za zamawiającego"])
    return d


# ------------------------------------------------------------------ 7. Checklista higieniczna
def checklista(przyklad):
    P = przyklad
    d = nowy_doc("Checklista instalacji higienicznej (przemysł spożywczy)", [
        (("Obiekt / odcinek", v(P, "Linia CIP L1-L3")), ("Data", v(P, "2026-11-19"))),
        (("Sprawdzający", v(P, "AB")), ("Wymagania klienta (dokument)", v(P, "Specyfikacja techniczna rew. 2"))),
    ], P)
    uwaga(d, "Odwołania ogólne: zasady higienicznego projektowania (np. wytyczne EHEDG), materiały do kontaktu z żywnością (rozporządzenie (WE) nr 1935/2004). Lista nie zastępuje specyfikacji zamawiającego.")
    pkt = ["Materiał zgodny z dokumentacją (gatunek), atesty kompletne", "Spadki zapewniają opróżnianie, brak zastoin",
           "Martwe strefy: odgałęzienia możliwie krótkie (wg wymagań klienta)", "Grań spoin: kolor w granicach wymagań, brak „cukru”",
           "Grań: brak wklęśnięć, nadlewów i niewspółosiowości ponad wymagania", "Chropowatość powierzchni wewnętrznej (jeśli wymagana)",
           "Złącza rozłączne i uszczelki - właściwe materiały, dostęp do mycia", "Podpory i mocowania - łatwe do mycia, bez szczelin, bez kontaktu z czarną stalą",
           "Mycie CIP - dostęp, odpowietrzenia, odwodnienia", "Powierzchnie po trawieniu i pasywacji - bez zacieków i rdzy",
           "Oznakowanie (medium, kierunek przepływu)", "Dokumentacja powykonawcza kompletna"]
    tabela(d, ["Lp.", "Punkt kontroli", "Spełnione", "Nie", "Nie dotyczy", "Uwagi"],
           [[str(i + 1), x, ("✔" if i != 10 else "") if P else "", ("✔" if i == 10 else "") if P else "", "", ("brak strzałek na L2" if i == 10 else "") if P else ""]
            for i, x in enumerate(pkt)], szer=[10, 72, 20, 14, 20, 44])
    podpisy(d, ["Sprawdzający", "Przedstawiciel zamawiającego"])
    return d


# ------------------------------------------------------------------ XLSX - pomocnicze
CIENKA = Side(style="thin", color="B9C3CE")
RAMKA = Border(left=CIENKA, right=CIENKA, top=CIENKA, bottom=CIENKA)
NAG = PatternFill("solid", fgColor="DCE8F3")
TYT = Font(bold=True, size=14, color="0B6FB8")


def arkusz_naglowek(ws, tytul, pola, przyklad, kolumn):
    ws["A1"] = tytul
    ws["A1"].font = TYT
    ws["A2"] = FIRMA_P + " · PRZYKŁAD WYPEŁNIENIA - dane fikcyjne" if przyklad else "[LOGO I DANE FIRMY - wpisz tutaj]"
    ws["A2"].font = Font(italic=True, size=9, color="5B6776")
    r = 3
    for k, val in pola:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 1).fill = NAG
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ws.cell(r, 4, val if przyklad else "")
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=min(kolumn, 9))
        for col in range(1, min(kolumn, 9) + 1):
            ws.cell(r, col).border = RAMKA
        r += 1
    return r + 1


def tabela_xlsx(ws, wiersz, naglowki, szer, dane, puste):
    for i, (n, s) in enumerate(zip(naglowki, szer), start=1):
        c = ws.cell(wiersz, i, n)
        c.font = Font(bold=True, size=9)
        c.fill = NAG
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = RAMKA
        ws.column_dimensions[get_column_letter(i)].width = s
    ws.row_dimensions[wiersz].height = 42
    for j, rzad in enumerate(dane + [[""] * len(naglowki)] * puste, start=wiersz + 1):
        for i in range(1, len(naglowki) + 1):
            c = ws.cell(j, i, rzad[i - 1] if i - 1 < len(rzad) else "")
            c.border = RAMKA
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(wiersz + 1, 1)
    ostatni = wiersz + len(dane) + puste
    ws.auto_filter.ref = f"A{wiersz}:{get_column_letter(len(naglowki))}{ostatni}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{wiersz}:{wiersz}"
    ws.oddFooter.center.text = WERSJA
    return ostatni


def stopka_xlsx(ws, wiersz, tekst):
    ws.cell(wiersz + 2, 1, tekst).font = Font(italic=True, size=8, color="5B6776")


# ------------------------------------------------------------------ 2. Dziennik
def dziennik(przyklad):
    P = przyklad
    wb = Workbook()
    ws = wb.active
    ws.title = "Dziennik spawania"
    r = arkusz_naglowek(ws, "Dziennik spawania", [("Zlecenie / obiekt", "Linia CIP - rozbudowa (fikcyjne)"), ("Zamawiający", "Mleczarnia Przykładowa (fikcyjna)"),
                                                  ("Nr rysunku / izometrii", "ISO-CIP-03 rew. B"), ("Okres prowadzenia", "2026-11-05 - 2026-11-19")], P, 15)
    nag = ["Lp.", "Data", "Nr złącza (wg rysunku)", "Wymiar (Ø × g)", "Materiał", "Nr WPS", "Spawacz (znak)", "Nr uprawnień / ważne do",
           "Nr wytopu / atestu materiału", "Partia spoiwa", "Formowanie (tak/nie)", "VT wynik", "Nr protokołu VT", "Inne badania (PT/RT, nr)", "Uwagi"]
    szer = [5, 11, 12, 11, 10, 12, 9, 18, 18, 11, 10, 8, 13, 14, 22]
    dane = [[1, "2026-11-05", "S-01", "51×2,0", "316L", "WPS-141-01", "AB", "[[nr]] / 2027-05-31", "W12345 / 3.1-778", "P-2201", "tak", "OK", "VT/2026/014", "-", ""],
            [2, "2026-11-05", "S-02", "51×2,0", "316L", "WPS-141-01", "AB", "[[nr]] / 2027-05-31", "W12345 / 3.1-778", "P-2201", "tak", "nOK", "VT/2026/014", "-", "wycięte, zob. S-02A"],
            [3, "2026-11-06", "S-02A", "51×2,0", "316L", "WPS-141-01", "AB", "[[nr]] / 2027-05-31", "W12345 / 3.1-778", "P-2201", "tak", "OK", "VT/2026/015", "-", "naprawa S-02"],
            [4, "2026-11-06", "S-03", "51×2,0", "316L", "WPS-141-01", "CD", "[[nr]] / 2026-12-15", "W12346 / 3.1-779", "P-2201", "tak", "OK", "VT/2026/014", "PT/2026/02", ""]] if P else []
    ost = tabela_xlsx(ws, r, nag, szer, dane, 0 if P else 40)
    dv = DataValidation(type="list", formula1='"tak,nie"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"K{r+1}:K{ost}")
    dv2 = DataValidation(type="list", formula1='"OK,nOK"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(f"L{r+1}:L{ost}")
    stopka_xlsx(ws, ost, "Wpisuj każde złącze. Wynik nOK → naprawa jako nowe złącze z dopiskiem (np. S-02A). Szablon nie zastępuje wymagań zamawiającego.")
    # zakładka spawaczy
    w2 = wb.create_sheet("Spawacze")
    w2["A1"] = "Lista spawaczy i ważność uprawnień"
    w2["A1"].font = TYT
    w2["A2"] = "Kolumna „Ważne do” zmienia kolor: czerwony = po terminie, żółty = mniej niż 60 dni."
    w2["A2"].font = Font(italic=True, size=9, color="5B6776")
    nag2 = ["Znak", "Imię i nazwisko", "Nr świadectwa / uprawnień", "Zakres (metoda, materiał, pozycje, grubość)", "Jednostka", "Ważne do", "Przedłużenie (data, podpis)"]
    dane2 = [["AB", "Adam Przykładowy (fikcyjny)", "[[nr]]", "141, 316L, PH, rura 2 mm [[wg świadectwa]]", "[[jednostka]]", "2027-05-31", ""],
             ["CD", "Celina Przykładowa (fikcyjna)", "[[nr]]", "141, 304L/316L, PC, rura 2-3 mm", "[[jednostka]]", "2026-12-15", ""]] if P else []
    ost2 = tabela_xlsx(w2, 4, nag2, [7, 26, 20, 36, 16, 12, 20], dane2, 0 if P else 20)
    rng = f"F5:F{ost2}"
    w2.conditional_formatting.add(rng, FormulaRule(formula=["AND(F5<>\"\",F5<TODAY())"], fill=PatternFill("solid", fgColor="F8CFCF")))
    w2.conditional_formatting.add(rng, FormulaRule(formula=["AND(F5<>\"\",F5-TODAY()<60)"], fill=PatternFill("solid", fgColor="FFF0B3")))
    if P:
        from datetime import date
        for row, dt in ((5, date(2027, 5, 31)), (6, date(2026, 12, 15))):
            w2.cell(row, 6, dt).number_format = "yyyy-mm-dd"
    return wb


# ------------------------------------------------------------------ 6. Rejestr atestów
def rejestr(przyklad):
    P = przyklad
    wb = Workbook()
    ws = wb.active
    ws.title = "Rejestr atestów"
    r = arkusz_naglowek(ws, "Rejestr atestów materiałowych", [("Zlecenie / obiekt", "Linia CIP - rozbudowa (fikcyjne)"), ("Zamawiający", "Mleczarnia Przykładowa (fikcyjna)")], P, 12)
    nag = ["Lp.", "Materiał (rura / kształtka / blacha / spoiwo)", "Gatunek", "Wymiar", "Producent / dostawca", "Nr wytopu", "Nr świadectwa", "Typ (2.2 / 3.1 / 3.2)",
           "Data świadectwa", "Gdzie użyto (złącza / odcinek)", "Link / lokalizacja skanu", "Uwagi"]
    dane = [[1, "rura", "1.4404 / 316L", "51×2,0", "Huta Przykładowa (fikcyjna)", "W12345", "3.1-778", "3.1", "2026-10-02", "S-01 - S-02A, L1", "atesty/3.1-778.pdf", ""],
            [2, "rura", "1.4404 / 316L", "51×2,0", "Huta Przykładowa (fikcyjna)", "W12346", "3.1-779", "3.1", "2026-10-02", "S-03 - S-10, L2", "atesty/3.1-779.pdf", ""],
            [3, "kolano 90°", "1.4404 / 316L", "51×2,0", "Dostawca Przykładowy", "K5521", "3.1-1204", "3.1", "2026-09-15", "L1, L2", "atesty/3.1-1204.pdf", ""],
            [4, "spoiwo (pręt)", "316L", "Ø 1,6", "Producent Przykładowy", "P-2201", "3.1-SP-88", "3.1", "2026-08-20", "wszystkie złącza", "atesty/SP-88.pdf", ""]] if P else []
    ost = tabela_xlsx(ws, r, nag, [5, 18, 14, 10, 22, 11, 13, 10, 12, 22, 22, 18], dane, 0 if P else 40)
    dv = DataValidation(type="list", formula1='"2.2,3.1,3.2"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"H{r+1}:H{ost}")
    stopka_xlsx(ws, ost, "Filtr w nagłówku: wybierz nr wytopu, żeby zobaczyć, gdzie użyto materiału. Typy świadectw wg PN-EN 10204 (odwołanie).")
    return wb


# ------------------------------------------------------------------ 8. Plan spawania i badań
def plan(przyklad):
    P = przyklad
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan spawania i badań"
    r = arkusz_naglowek(ws, "Plan spawania i badań", [("Zlecenie / obiekt", "Linia CIP - rozbudowa (fikcyjne)"), ("Zamawiający", "Mleczarnia Przykładowa (fikcyjna)"),
                                                     ("Wymagania / normy klienta", "Specyfikacja techn. rew. 2; VT wg PN-EN ISO 17637; kryteria PN-EN ISO 5817 poz. B"),
                                                     ("Data i rewizja planu", "2026-11-03 / rew. 0")], P, 11)
    nag = ["Lp.", "Grupa złączy / odcinek", "Liczba złączy", "Materiał / wymiar", "Nr WPS", "Wymagane kwalifikacje spawaczy", "Badania i zakres (VT / PT / RT, %)",
           "Kryteria akceptacji", "Kto bada", "Termin", "Uwagi"]
    dane = [[1, "L1 - zasilanie CIP", 18, "316L, Ø51×2", "WPS-141-01", "141, rura, PH/PC", "VT 100%, endoskop grani 100%", "ISO 5817 B; grań poz. ≤ 2", "wykonawca / [[jednostka]]", "tydz. 46", ""],
            [2, "L2 - powrót CIP", 22, "316L, Ø51×2", "WPS-141-01", "141, rura, PH/PC", "VT 100%, PT 10%", "ISO 5817 B", "wykonawca / [[jednostka]]", "tydz. 46-47", "[[wymagania klienta]]"],
            [3, "L3 - zbiornik", 9, "316L, Ø76×2", "WPS-141-02", "141, rura, PH", "VT 100%", "ISO 5817 B", "wykonawca", "tydz. 47", ""]] if P else []
    ost = tabela_xlsx(ws, r, nag, [5, 22, 9, 16, 12, 20, 24, 20, 18, 11, 20], dane, 0 if P else 20)
    stopka_xlsx(ws, ost, "Plan przekaż zamawiającemu przed rozpoczęciem prac, jeśli tego wymaga. Zakres badań ustala zamawiający / projekt.")
    return wb


DOKUMENTY = [
    ("01-WPS-TIG-141.docx", wps),
    ("02-Dziennik-spawania.xlsx", dziennik),
    ("03-Protokol-VT.docx", vt),
    ("04-Protokol-proby-szczelnosci.docx", proba),
    ("05-Protokol-odbioru.docx", odbior),
    ("06-Rejestr-atestow.xlsx", rejestr),
    ("07-Checklista-instalacji-higienicznej.docx", checklista),
    ("08-Plan-spawania-i-badan.xlsx", plan),
]


def main():
    for folder, przyklad in (("puste", False), ("przyklady", True)):
        os.makedirs(os.path.join(OUT, folder), exist_ok=True)
        for nazwa, fn in DOKUMENTY:
            plik = os.path.join(OUT, folder, nazwa if not przyklad else nazwa.replace(".", "-PRZYKLAD.", 1))
            fn(przyklad).save(plik)
            print(plik)


if __name__ == "__main__":
    main()
